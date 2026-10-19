"""Kelime listesini (xlsx/txt/csv) okur, veri raporu ve tohum dosyaları üretir.

Geliştirme zamanında tools/build_seed.py bu modülü kullanarak seed/ klasörünü
oluşturur; uygulama ilk açılışta seed/pool.json'u veritabanına yükler.
"""
from __future__ import annotations

import collections
import csv
import json
import re
from pathlib import Path
from typing import Any

from . import textutil as tu

MAIN_SHEET = "Havuz (2.200 Kart)"
PV_SHEET = "Öbek Fiiller (110)"
RESERVE_SHEET = "Yedek (1.300 Kart)"

# xlsx "Başla Buradan" sayfasındaki kapsama eğrisi (ÖSYM tam sınav metinlerinde tanıma oranı)
COVERAGE_CURVE = [(0, 92.1), (500, 94.2), (1000, 95.5), (1500, 96.3), (2000, 96.8), (2200, 96.9),
                  (2500, 97.0), (3000, 97.2)]


def _rows(ws) -> tuple[list[str], list[dict]]:
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(it)]
    data = []
    for r in it:
        if r is None or not any(v is not None for v in r):
            continue
        data.append(dict(zip(hdr, r)))
    return hdr, data


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        if v in ("", "–", "-"):
            return None
    return v


def read_pool_xlsx(path: Path) -> dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out: dict[str, Any] = {"sheets": {}, "items": []}
    names = wb.sheetnames
    for name in names:
        ws = wb[name]
        out["sheets"][name] = {"rows": ws.max_row, "cols": ws.max_column}

    def get(name):
        return _rows(wb[name]) if name in names else ([], [])

    hdr_main, main = get(MAIN_SHEET)
    hdr_pv, pv = get(PV_SHEET)
    hdr_res, res = get(RESERVE_SHEET)
    out["headers"] = {MAIN_SHEET: hdr_main, PV_SHEET: hdr_pv, RESERVE_SHEET: hdr_res}
    out["raw"] = {MAIN_SHEET: main, PV_SHEET: pv, RESERVE_SHEET: res}

    items = []
    for d in main:
        word = _clean(d.get("Kelime"))
        if not word:
            continue
        kind = "compound" if ("-" in word or " " in word) else "word"
        items.append({
            "lemma": tu.norm_en(word), "display": word.strip(), "kind": kind, "pool": "main",
            "pos": _clean(d.get("Tür")), "meaning": _clean(d.get("Türkçe Anlam")),
            "family": _clean(d.get("Aile (aynı kartta)")), "level": _clean(d.get("Seviye")),
            "text_prob": _clean(d.get("Metinde Çıkma Olasılığı (%)")),
            "option_prob": _clean(d.get("Şık Olarak Çıkma (her sınav, %)")),
            "src_count": _clean(d.get("YDS Kaynak Sayısı (/12)")),
            "osym_evidence": _clean(d.get("ÖSYM Kanıtı")), "tags": _clean(d.get("Etiket")),
            "sources": _clean(d.get("YDS Kaynakları")), "priority": _clean(d.get("Öncelik Puanı (0–100)")),
            "rank": _clean(d.get("Sıra")), "plan_week": _clean(d.get("Hafta")), "plan_day": _clean(d.get("Gün")),
        })
    for d in pv:
        word = _clean(d.get("Öbek Fiil"))
        if not word:
            continue
        asked = _clean(d.get("ÖSYM'de Şık Olarak Sorulma")) or 0
        items.append({
            "lemma": tu.norm_en(word), "display": word.strip(), "kind": "phrasal", "pool": "main",
            "pos": "fiil", "meaning": _clean(d.get("Türkçe Anlam")), "family": None, "level": None,
            "text_prob": None, "option_prob": None, "src_count": _clean(d.get("YDS Kaynak Sayısı")),
            "osym_evidence": (f"ÖSYM'de {asked} kez şık olarak soruldu" if asked else None),
            "tags": ("uzman seçimi" if _clean(d.get("Uzman Seçimi")) == "evet" else None),
            "sources": None, "priority": None, "osym_asked": asked,
            "rank": _clean(d.get("Sıra")), "plan_week": _clean(d.get("Hafta")), "plan_day": _clean(d.get("Gün")),
        })
    for d in res:
        word = _clean(d.get("Kelime"))
        if not word:
            continue
        kind = "compound" if ("-" in word or " " in word) else "word"
        items.append({
            "lemma": tu.norm_en(word), "display": word.strip(), "kind": kind, "pool": "reserve",
            "pos": _clean(d.get("Tür")), "meaning": _clean(d.get("Türkçe Anlam")),
            "family": _clean(d.get("Aile")), "level": _clean(d.get("Seviye")),
            "text_prob": _clean(d.get("Metinde Çıkma Olasılığı (%)")),
            "option_prob": _clean(d.get("Şık Olarak Çıkma (her sınav, %)")),
            "src_count": _clean(d.get("YDS Kaynak Sayısı (/12)")),
            "osym_evidence": _clean(d.get("ÖSYM Kanıtı")), "tags": _clean(d.get("Etiket")),
            "sources": None, "priority": _clean(d.get("Öncelik Puanı (0–100)")),
            "rank": _clean(d.get("Sıra")), "plan_week": None, "plan_day": None,
        })
    out["items"] = items
    return out


def assign_order(items: list[dict]) -> None:
    """Öğrenme sırası: ana havuz Sıra'ya göre; öbek fiiller planlandıkları günün sonuna."""
    main = [i for i in items if i["pool"] == "main" and i["kind"] != "phrasal"]
    pv = [i for i in items if i["kind"] == "phrasal"]
    res = [i for i in items if i["pool"] == "reserve"]
    main.sort(key=lambda i: i["rank"] or 1e9)
    day_last: dict[int, float] = {}
    for idx, it in enumerate(main):
        it["order_index"] = float(idx + 1)
        if it.get("plan_day"):
            day_last[int(it["plan_day"])] = float(idx + 1)
    for it in pv:
        # 'Gün' iki sayfada da genel gün numarası (1-42)
        day = it.get("plan_day")
        base = day_last.get(int(day)) if day else None
        if base is None:
            base = float(len(main))
        it["order_index"] = base + 0.01 * float(it["rank"] or 1)
    res.sort(key=lambda i: i["rank"] or 1e9)
    start = float(len(main) + len(pv) + 100)
    for idx, it in enumerate(res):
        it["order_index"] = start + idx


def coverage_weights(items: list[dict]) -> None:
    """Her kartın sınav metni kapsamasına (yüzde puan) katkısı — puan tahmin modeli için."""
    main = sorted([i for i in items if i["pool"] == "main" and i["kind"] != "phrasal"], key=lambda i: i["order_index"])
    res = sorted([i for i in items if i["pool"] == "reserve"], key=lambda i: i["order_index"])
    seq = main + res
    pts = COVERAGE_CURVE
    for (n0, c0), (n1, c1) in zip(pts[:-1], pts[1:]):
        block = seq[n0:n1]
        if not block:
            continue
        gain = c1 - c0
        tp = [max(float(i.get("text_prob") or 1.0), 0.5) for i in block]
        s = sum(tp)
        for it, t in zip(block, tp):
            it["cov_weight"] = round(gain * t / s, 6)
    for it in items:
        if it["kind"] == "phrasal":
            it["cov_weight"] = 0.004
        it.setdefault("cov_weight", 0.0005)


def add_zipf(items: list[dict]) -> None:
    try:
        from wordfreq import zipf_frequency
    except Exception:
        return
    for it in items:
        it["zipf"] = round(zipf_frequency(it["lemma"], "en"), 2)


# ---------------------------------------------------------------- karışabilir çiftler
def _bigrams(w: str) -> set[str]:
    w = f"^{w}$"
    return {w[i:i + 2] for i in range(len(w) - 1)}


def _tr_roots_shared(m1: str | None, m2: str | None, n: int = 5) -> bool:
    t1 = [t for t in tu.meaning_tokens(m1) if " " not in t and len(t) >= n]
    t2 = [t for t in tu.meaning_tokens(m2) if " " not in t and len(t) >= n]
    return any(tu.common_prefix(a, b) >= n for a in t1 for b in t2)


def find_confusables(items: list[dict]) -> list[dict]:
    """Yazılışı benzer (OSA mesafesi) veya anlamı çakışan kelime çiftleri.

    kind: 'yazılış' | 'anlam' | 'yazılış+anlam' | 'aile' (aynı kökün biçimleri; paket engeli yok)
    """
    pool = [i for i in items if i["pool"] == "main"]
    by_lemma = {i["lemma"]: i for i in pool}
    lemmas = [i["lemma"] for i in pool]
    index: dict[str, list[int]] = collections.defaultdict(list)
    grams = []
    for k, w in enumerate(lemmas):
        g = _bigrams(w)
        grams.append(g)
        for b in g:
            index[b].append(k)
    pairs: dict[tuple[str, str], dict] = {}
    for k, w in enumerate(lemmas):
        if len(w) < 4:
            continue
        counts: collections.Counter = collections.Counter()
        for b in grams[k]:
            for j in index[b]:
                if j > k:
                    counts[j] += 1
        for j, c in counts.items():
            v = lemmas[j]
            if len(v) < 4 or abs(len(v) - len(w)) > 2:
                continue
            jac = c / len(grams[k] | grams[j])
            if jac < 0.45:
                continue
            d = tu.osa_distance(w, v, max_d=2)
            L = max(len(w), len(v))
            # aynı ailenin biçimleri (significant/significantly) karışıklık sayılmaz
            if w.startswith(v) or v.startswith(w):
                continue
            ok = (d == 1 and L >= 4) or (d == 2 and L >= 6)
            if not ok:
                # öneki farklı aynı kök: precede/proceed, prescribe/proscribe, emigrate/immigrate
                if d <= 3 and L >= 7 and w[-4:] == v[-4:] and tu.common_prefix(w, v) >= 1:
                    ok = d <= 2 or (len(w) == len(v) and d <= 3)
            if ok:
                key = (w, v) if w < v else (v, w)
                if w.replace(" ", "").replace("-", "") == v.replace(" ", "").replace("-", ""):
                    kind = "aile"  # breakdown / break down
                elif (tu.common_prefix(w, v) >= max(5, min(len(w), len(v)) - 3)
                      and _tr_roots_shared(by_lemma[w].get("meaning"), by_lemma[v].get("meaning"))):
                    kind = "aile"  # dependence / dependency, diversity / diversify
                else:
                    kind = "yazılış"
                pairs[key] = {"a": key[0], "b": key[1], "kind": kind, "score": round(1 - d / L, 3)}
    # anlam çakışması: aynı Türkçe ifadeyi paylaşanlar (ifade ne kadar az kelimede geçiyorsa o kadar özgül)
    by_tok: dict[str, list[int]] = collections.defaultdict(list)
    for k, it in enumerate(pool):
        for t in tu.meaning_phrases(it.get("meaning")):
            if len(t) >= 4:
                by_tok[t].append(k)
    for t, ks in by_tok.items():
        if len(ks) < 2 or len(ks) > 8:
            continue
        spec = round(1.0 / (len(ks) - 1), 3)
        for x in range(len(ks)):
            for y in range(x + 1, len(ks)):
                a, b = pool[ks[x]], pool[ks[y]]
                if a["lemma"].startswith(b["lemma"]) or b["lemma"].startswith(a["lemma"]):
                    continue
                key = (a["lemma"], b["lemma"]) if a["lemma"] < b["lemma"] else (b["lemma"], a["lemma"])
                if key in pairs:
                    if pairs[key]["kind"] == "yazılış":
                        pairs[key]["kind"] = "yazılış+anlam"
                    if pairs[key]["kind"] == "anlam":
                        pairs[key]["shared"].append(t)
                        pairs[key]["score"] = round(min(1.0, pairs[key]["score"] + spec / 2), 3)
                    continue
                pairs[key] = {"a": key[0], "b": key[1], "kind": "anlam", "score": spec, "shared": [t]}
    return list(pairs.values())


# ---------------------------------------------------------------- çevrimdışı sözlük
def build_offline_dict(sources_dir: Path, items: list[dict]) -> dict[str, dict]:
    """Hızlı ekleme için çevrimdışı İngilizce→Türkçe sözlük (kaynak klasörlerinden)."""
    d: dict[str, dict] = {}

    def put(w: str, tr: str, pos: str | None, src: str, overwrite: bool = False):
        w = tu.norm_en(w)
        if not w or not tr:
            return
        if w in d and not overwrite:
            return
        d[w] = {"tr": tr.strip(), "pos": (pos or "").strip() or None, "src": src}

    turk = sources_dir / "02_Turk_YDS_Kaynaklari"
    js = turk / "Turksev_github_YDS_kelime-dizin_9379_kelime.js"
    if js.exists():
        txt = js.read_text(encoding="utf-8")
        for m in re.finditer(r'\{e:"((?:[^"\\]|\\.)*)",t:"((?:[^"\\]|\\.)*)"(?:,p:[\d.]+)?(?:,k:\d+)?(?:,y:"((?:[^"\\]|\\.)*)")?\}', txt):
            put(m.group(1), m.group(2).replace("…", "").strip(" ;,"), m.group(3), "turksev")
    afl = turk / "EnglishAFL_Kayra_YDT-YDS_Kelime_Listesi_3239.txt"
    if afl.exists():
        for line in afl.read_text(encoding="utf-8", errors="ignore").splitlines():
            if ":" in line and not line.startswith("---"):
                w, tr = line.split(":", 1)
                if re.fullmatch(r"[A-Za-z][A-Za-z' \-]*", w.strip()):
                    put(w, tr, None, "afl")
    seyfi = turk / "MetehanGecmen_SeyfiHoca_YDS_Cikmis_600_Kelime.csv"
    if seyfi.exists():
        with seyfi.open(encoding="utf-8", errors="ignore") as f:
            for row in csv.DictReader(f):
                put(row.get("English", ""), row.get("Turkish", ""), None, "seyfi")
    # listenin kendi anlamları en yüksek öncelik
    for it in items:
        if it.get("meaning"):
            put(it["lemma"], it["meaning"], it.get("pos"), "liste", overwrite=True)
    return d


# ---------------------------------------------------------------- gerçek ÖSYM soruları
def read_osym_questions(analysis_xlsx: Path) -> list[dict]:
    if not analysis_xlsx.exists():
        return []
    import openpyxl
    wb = openpyxl.load_workbook(str(analysis_xlsx), read_only=True, data_only=True)
    if "Tüm Sorular" not in wb.sheetnames:
        return []
    hdr, rows = _rows(wb["Tüm Sorular"])
    out = []
    for r in rows:
        typ = r.get("Soru Tipi")
        if typ not in ("Kelime Bilgisi", "Dil Bilgisi"):
            continue
        opts = [r.get(k) for k in "ABCDE"]
        if not all(opts) or not r.get("Soru Kökü"):
            continue
        out.append({
            "session": r.get("Oturum"), "no": r.get("Soru No"), "type": typ, "sub": r.get("Alt Konu"),
            "stem": str(r.get("Soru Kökü")).strip(), "options": [str(o).strip() for o in opts],
            "answer": r.get("Doğru Cevap"), "answer_text": r.get("Doğru Cevabın Metni"),
            "note": r.get("Not / İpucu"),
        })
    return out


def read_linker_heatmap(analysis_xlsx: Path) -> list[dict]:
    """Analiz dosyasındaki bağlaç/geçiş ifadeleri ısı haritası (öneri listesi için)."""
    if not analysis_xlsx.exists():
        return []
    import openpyxl
    wb = openpyxl.load_workbook(str(analysis_xlsx), read_only=True, data_only=True)
    name = "Isı Haritası-Bağlaçlar"
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    hdr_idx = None
    for i, r in enumerate(rows):
        if r and r[0] == "İfade":
            hdr_idx = i
            break
    if hdr_idx is None:
        return []
    hdr = rows[hdr_idx]
    for r in rows[hdr_idx + 1:]:
        if not r or not r[0]:
            continue
        d = dict(zip(hdr, r))
        out.append({"phrase": str(d["İfade"]).strip(), "relation": d.get("İlişki"), "total": d.get("Toplam"),
                    "in_options": d.get("Şıklarda"), "correct": d.get("Doğru Şıkta")})
    return out


# ---------------------------------------------------------------- veri raporu
def data_report(xlsx_path: Path, txt_path: Path | None, parsed: dict, confus: list[dict]) -> str:
    items = parsed["items"]
    main = [i for i in items if i["pool"] == "main" and i["kind"] != "phrasal"]
    pv = [i for i in items if i["kind"] == "phrasal"]
    res = [i for i in items if i["pool"] == "reserve"]
    L: list[str] = []
    w = L.append
    w("# Veri Raporu — YDS Akıllı Kelime Havuzu\n")
    w(f"Kaynak dosya: `{xlsx_path}`\n")
    w("## 1. Özet sayılar\n")
    w("| Bölüm | Kayıt | Not |")
    w("|---|---:|---|")
    w(f"| Ana havuz (`{MAIN_SHEET}`) | {len(main)} | kelime kartı; {sum(1 for i in main if i.get('family'))} kartta aile biçimi var |")
    w(f"| Öbek fiiller (`{PV_SHEET}`) | {len(pv)} | phrasal verb |")
    w(f"| **Çalışma listesi toplamı** | **{len(main) + len(pv)}** | uygulamanın hedef listesi |")
    w(f"| Yedek (`{RESERVE_SHEET}`) | {len(res)} | liste biterse/zaman kalırsa |")
    fam_forms = sum(len([f for f in re.split(r'[,;]\s*', i['family']) if f.strip()]) for i in main if i.get("family"))
    w(f"\nAile biçimleriyle ana havuzdaki farklı kelime sayısı: {len(main) + fam_forms}.\n")
    w("Çalışma kitabındaki sayfalar: " + ", ".join(f"`{k}` ({v['rows']}×{v['cols']})" for k, v in parsed["sheets"].items()) + "\n")

    w("## 2. Kolonlar\n")
    for name, hdr in parsed["headers"].items():
        if hdr:
            w(f"- **{name}**: " + ", ".join(f"`{h}`" for h in hdr))
    w("")

    w("## 3. Eksik değerler\n")
    for name, rows in parsed["raw"].items():
        hdr = parsed["headers"][name]
        if not rows:
            continue
        miss = []
        for h in hdr:
            n = sum(1 for r in rows if _clean(r.get(h)) is None)
            if n:
                miss.append(f"`{h}`: {n}")
        w(f"- **{name}** ({len(rows)} satır): " + (", ".join(miss) if miss else "eksik yok"))
    w("\nYorum: Kelime, tür ve Türkçe anlam kolonlarında **hiç eksik yok**. `Seviye` boşlukları (CEFR listelerinde bulunmayan"
      " kelimeler) ve `ÖSYM Kanıtı`/`Etiket` boşlukları doğaldır; uygulamada bu alanlar zorunlu değil.\n")

    w("## 4. Tekrarlı kayıtlar\n")
    cnt = collections.Counter(i["lemma"] for i in items)
    dups = [k for k, v in cnt.items() if v > 1]
    w(f"- Sayfalar arası/sayfa içi aynı kelime: **{len(dups)}** " + (f"({', '.join(dups[:20])})" if dups else "(yok)"))
    fam_map = collections.defaultdict(list)
    heads = {i["lemma"] for i in items}
    for i in items:
        if i.get("family"):
            for f in re.split(r"[,;]\s*", i["family"]):
                f = tu.norm_en(f)
                if f:
                    fam_map[f].append(i["lemma"])
    fam_head = [f for f in fam_map if f in heads]
    fam_multi = [f for f, v in fam_map.items() if len(v) > 1]
    w(f"- Aile biçimi olup aynı zamanda ayrı kart olan kelime: **{len(fam_head)}**" + (f" ({', '.join(fam_head[:15])})" if fam_head else ""))
    w(f"- Birden fazla kartta aile biçimi olarak geçen kelime: **{len(fam_multi)}**")
    stem_groups = collections.defaultdict(list)
    for i in main:
        stem_groups[i["lemma"][:6]].append(i["lemma"])
    near_dup = [(a, b) for a in heads for b in heads if a < b and (a + "ly" == b or a + "s" == b)]
    w(f"- Yakın tekrar (ör. `x` ve `xly` ayrı kart): **{len(near_dup)}**" + (f" — {', '.join(a + '/' + b for a, b in near_dup[:12])}" if near_dup else ""))
    w("")

    w("## 5. Çok anlamlı kelimeler\n")
    ng = collections.Counter(len(tu.meaning_groups(i.get("meaning"))) for i in main)
    w("Türkçe anlamda `;` ile ayrılan anlam grubu sayısı (ana havuz):\n")
    w("| Anlam grubu | Kart |")
    w("|---:|---:|")
    for k in sorted(ng):
        w(f"| {k} | {ng[k]} |")
    multi_pos = [i for i in main if len(tu.pos_list(i.get("pos"))) > 1]
    w(f"\n- Birden fazla anlam grubu olan kart: **{sum(v for k, v in ng.items() if k >= 2)}**")
    w(f"- Birden fazla sözcük türü olan kart (ör. `isim/fiil`): **{len(multi_pos)}**")
    ex3 = [i for i in main if len(tu.meaning_groups(i.get("meaning"))) >= 3][:12]
    w("- 3+ anlam grubu örnekleri: " + "; ".join(f"**{i['display']}** ({i['meaning']})" for i in ex3))
    w("\nUygulamada: listedeki anlam asla ezilmez. Çoktan seçmeli şıklarda ilk anlam grubu kısaltılarak gösterilir;"
      " kart ekranında ve ayrıntıda tüm anlamlar görünür. YDS'deki kullanım anlamları ayrıca 'doğrulanmamış' içerik olarak eklenir.\n")

    w("## 6. Phrasal verb'ler ve çok kelimeli ifadeler\n")
    w(f"- Öbek fiil sayfasında **{len(pv)}** phrasal verb var; ÖSYM'de şık olarak sorulmuş olanlar: "
      f"**{sum(1 for i in pv if (i.get('osym_asked') or 0) > 0)}**.")
    w("- İlk 25: " + ", ".join(i["display"] for i in sorted(pv, key=lambda i: i.get('rank') or 999)[:25]))
    comp = [i["display"] for i in main if i["kind"] == "compound"]
    w(f"- Ana havuzda tireli/çok kelimeli kart: **{len(comp)}** ({', '.join(comp)})")
    w("")

    w("## 7. Dağılımlar\n")
    w("- Tür: " + ", ".join(f"{k}: {v}" for k, v in collections.Counter(i.get('pos') for i in main).most_common()))
    w("- Seviye: " + ", ".join(f"{k or 'yok'}: {v}" for k, v in collections.Counter(i.get('level') for i in main).most_common()))
    w("- Etiket: " + ", ".join(f"{k}: {v}" for k, v in collections.Counter(i.get('tags') for i in main if i.get('tags')).most_common()))
    lens = [len(i["lemma"]) for i in main]
    w(f"- Kelime uzunluğu: en kısa {min(lens)}, en uzun {max(lens)}, ortalama {sum(lens) / len(lens):.1f} harf")
    zs = sorted(i.get("zipf") or 0 for i in main)
    if zs and zs[-1] > 0:
        w(f"- wordfreq Zipf sıklığı: medyan {zs[len(zs) // 2]:.2f}, %10 {zs[len(zs) // 10]:.2f}, %90 {zs[len(zs) * 9 // 10]:.2f}")
    w("")

    w("## 8. Karışabilir kelime çiftleri (otomatik tespit)\n")
    kinds = collections.Counter(c["kind"] for c in confus)
    w("Aynı pakette ilk kez öğretilmeyecek çiftler: " + ", ".join(f"{k}: {v}" for k, v in kinds.items()))
    ort = [c for c in confus if "yazılış" in c["kind"]]
    ort.sort(key=lambda c: -c["score"])
    w("\nYazılışı benzer örnekler: " + ", ".join(f"{c['a']}/{c['b']}" for c in ort[:40]))
    sem = [c for c in confus if c["kind"] == "anlam"]
    sem.sort(key=lambda c: -c["score"])
    w("\nAnlamı çakışan (en özgül) örnekler: " + ", ".join(f"{c['a']}/{c['b']} ({c['shared'][0]})" for c in sem[:30]))
    fam = [c for c in confus if c["kind"] == "aile"]
    w(f"\nAynı kökten biçimler ('aile', paket engeli yok, ileride biçim ayırt etme alıştırması): {len(fam)} — "
      + ", ".join(f"{c['a']}/{c['b']}" for c in fam[:15]))
    w("")

    if txt_path and txt_path.exists():
        lines = txt_path.read_text(encoding="utf-8").splitlines()
        tw = {tu.norm_en(re.sub(r"\s*\(.*\)$", "", l.split("\t")[0])) for l in lines if l.strip()}
        lw = {i["lemma"] for i in main + pv}
        w("## 9. Anki/Quizlet metin dosyası tutarlılığı\n")
        w(f"- `{txt_path.name}`: {len(lines)} satır; xlsx çalışma listesiyle ortak: {len(tw & lw)}; "
          f"yalnız txt'de: {len(tw - lw)}; yalnız xlsx'te: {len(lw - tw)}.")
        w("")

    w("## 10. Uygulamanın veriyi kullanma biçimi\n")
    w("- Öğrenme sırası: ana havuz `Sıra` (öncelik) sırasıyla; öbek fiiller xlsx'teki gün planına göre araya yerleşir.")
    w("- `Öncelik Puanı`, `Metinde Çıkma` ve `Şık Olarak Çıkma` olasılıkları: yeni kelime sırası, sınav-günü kapsama ve puan tahmini.")
    w("- Kapsama eğrisi (xlsx 'Başla Buradan'): " + ", ".join(f"{n} kart→%{c}" for n, c in COVERAGE_CURVE) + ".")
    w("- Yedek 1.300 kart varsayılan olarak plana girmez; Ayarlar'dan açılabilir.")
    return "\n".join(L) + "\n"
